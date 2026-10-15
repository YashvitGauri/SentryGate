"""Shared look of the college progress reports (Report-1 format).

Fonts, headings, bullets, tables, figures, page setup and the cover page.
Each report script (build_progress2.py, build_progress3.py) only writes its
own words and numbers, so the reports always look the same.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = HERE / "out"

# Names and enrollment numbers live in reports/team.json, which git ignores
# because the repository is public. Copy team.example.json to start.
_TEAM_FILE = HERE / "team.json"
if not _TEAM_FILE.exists():
    raise SystemExit("Missing reports/team.json - copy reports/team.example.json and fill it in.")
TEAM = json.loads(_TEAM_FILE.read_text())

SERIF = "Times New Roman"
SANS = "Arial"
HEADER_FILL = "DCE6F2"          # light blue-grey header row, as in Report-1
GREEN = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xC0, 0x6A, 0x00)
RED = RGBColor(0xC6, 0x28, 0x28)
GREY = RGBColor(0x55, 0x55, 0x55)



def pct(x: float) -> str:
    return f"{x:.1f}%"


# ---------------------------------------------------------------- helpers

def set_font(run, name=SERIF, size=12, bold=None, italic=None, color=None, underline=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if color is not None:
        run.font.color.rgb = color


def rich(par, text: str, size=12, name=SERIF, color=None):
    """Add text with **bold** spans."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            set_font(par.add_run(chunk), name=name, size=size, bold=(i % 2 == 1), color=color)
    return par


def para(doc, text: str, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, name=SERIF):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    return rich(p, text, size=size, name=name)


def heading(doc, text: str, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=size, bold=True, underline=True)
    return p


def bullet(doc, lead: str, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(lead), bold=True)
    rich(p, f" - {text}")


def caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = False
    set_font(p.add_run(text), size=9, italic=True)


def shade(cell, fill: str):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc, header: list[str], rows: list[list], widths: list[float], center_cols=(), bold_row=None):
    """rows: list of lists; a cell may be str or (str, RGBColor) for coloured bold text."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, HEADER_FILL)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
        set_font(c.paragraphs[0].add_run(h), size=11, bold=True)
    for r, row in enumerate(rows, start=1):
        for i, val in enumerate(row):
            c = t.rows[r].cells[i]
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            if isinstance(val, tuple):
                set_font(p.add_run(val[0]), size=11, bold=True, color=val[1])
            else:
                set_font(p.add_run(str(val)), size=11, bold=(bold_row == r - 1))
    # Fixed column widths. Word and LibreOffice read the grid, not just the
    # cells, so both are set; without this every column comes out equal.
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    style_table(t)
    return t


def style_table(t):
    """Report-1 look: light grey borders, padded cells, centred vertically,
    and the table kept on one page with its header."""
    tbl_pr = t._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 45), ("bottom", 45), ("left", 100), ("right", 100)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.append(margins)
    for r, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if r == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            tc_pr = cell._element.get_or_add_tcPr()
            v = OxmlElement("w:vAlign")
            v.set(qn("w:val"), "center")
            tc_pr.append(v)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = True  # table stays with its caption


def figure(doc, path: Path, width_in: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_in))


def new_report(title: str) -> Document:
    """A4 document with the Report-1 margins, font and cover page."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.3))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    base = doc.styles["Normal"]
    base.font.name = SERIF
    base.font.size = Pt(12)

    # ---- cover -----------------------------------------------------------
    def cline(text, size=12, bold=False, underline=False, before=0, after=4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        set_font(p.add_run(text), name=SANS, size=size, bold=bold, underline=underline)

    cline("SentryGate", 20, bold=True, underline=True, before=30, after=14)
    cline("A Security-Aware Semantic Caching Gateway for Large Language", 14, bold=True, after=0)
    cline("Model Applications", 14, bold=True, after=18)
    cline(title, 12.5, bold=True, underline=True, after=18)
    cline("Submitted in the partial fulfilment of the Degree of", 12, after=2)
    cline("Bachelor of Technology", 12, bold=True, after=2)
    cline("(Information Technology)", 12, after=16)
    cline("Submitted by", 12, bold=True, after=6)
    for m in TEAM["members"]:
        cline(m, 12, after=2)
    cline("Under the supervision of", 12, bold=True, before=14, after=6)
    cline(TEAM["supervisor"], 13, bold=True, after=10)
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(str(HERE / "assets/msit_logo.png"), width=Inches(1.45))
    cline("Department of Information Technology", 12, bold=True, before=8, after=2)
    cline("Maharaja Surajmal Institute of Technology", 12, bold=True, after=2)
    cline("Janakpuri, New Delhi.", 12, after=2)
    cline("2023-27", 12, after=0)
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    return doc


def page_break(doc) -> None:
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def save(doc, name: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / name
    doc.save(out)
    return out
